"""Навыки создания документов и презентаций: docx / pptx / pdf / xlsx.

Зависимости (python-docx, python-pptx, reportlab, openpyxl) — опциональные
(extras `docs`). Если библиотека не установлена, навык возвращает понятную
ошибку вместо падения агента.
"""

from __future__ import annotations

from .base import Skill, SkillContext, SkillParam, SkillResult, SkillSpec

DOCUMENTS_SUBDIR = "documents"


def _ensure_suffix(filename: str, suffix: str) -> str:
    return filename if filename.lower().endswith(suffix) else filename + suffix


class CreateDocxSkill(Skill):
    spec = SkillSpec(
        name="documents.create_docx",
        description="Создать документ Word (.docx) с заголовком и текстом (абзацы разделяются пустой строкой).",
        parameters=[
            SkillParam("filename", "string", "Имя файла, например report.docx"),
            SkillParam("title", "string", "Заголовок документа"),
            SkillParam("content", "string", "Текст документа"),
        ],
    )

    def _run(self, context: SkillContext, filename: str, title: str, content: str) -> SkillResult:
        try:
            from docx import Document
        except ImportError:
            return SkillResult(ok=False, error="python-docx не установлен (extras: docs)")

        target = context.resolve_path(_ensure_suffix(filename, ".docx"), subdir=DOCUMENTS_SUBDIR)
        context.policy.enforce("files.write", path=target)
        target.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()
        doc.add_heading(title, level=0)
        for paragraph in content.split("\n\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
        doc.save(target)
        return SkillResult(ok=True, output=f"Документ создан: {target}", data={"path": str(target)})


class CreatePptxSkill(Skill):
    spec = SkillSpec(
        name="documents.create_pptx",
        description="Создать презентацию PowerPoint (.pptx) из списка слайдов с заголовком и маркированным списком.",
        parameters=[
            SkillParam("filename", "string", "Имя файла, например deck.pptx"),
            SkillParam("title", "string", "Заголовок презентации (титульный слайд)"),
            SkillParam(
                "slides",
                "array",
                "Список слайдов: [{\"title\": str, \"bullets\": [str, ...]}, ...]",
                items={"type": "object"},
            ),
        ],
    )

    def _run(self, context: SkillContext, filename: str, title: str, slides: list[dict]) -> SkillResult:
        try:
            from pptx import Presentation
        except ImportError:
            return SkillResult(ok=False, error="python-pptx не установлен (extras: docs)")

        target = context.resolve_path(_ensure_suffix(filename, ".pptx"), subdir=DOCUMENTS_SUBDIR)
        context.policy.enforce("files.write", path=target)
        target.parent.mkdir(parents=True, exist_ok=True)

        prs = Presentation()
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = title

        bullet_layout = prs.slide_layouts[1]
        for slide_data in slides:
            slide = prs.slides.add_slide(bullet_layout)
            slide.shapes.title.text = str(slide_data.get("title", ""))
            body = slide.placeholders[1].text_frame
            bullets = slide_data.get("bullets", [])
            if bullets:
                body.text = str(bullets[0])
                for bullet in bullets[1:]:
                    p = body.add_paragraph()
                    p.text = str(bullet)

        prs.save(target)
        return SkillResult(
            ok=True,
            output=f"Презентация создана: {target} ({len(slides)} слайдов + титульный)",
            data={"path": str(target)},
        )


class CreatePdfSkill(Skill):
    spec = SkillSpec(
        name="documents.create_pdf",
        description="Создать PDF-документ с заголовком и текстом (абзацы разделяются пустой строкой).",
        parameters=[
            SkillParam("filename", "string", "Имя файла, например report.pdf"),
            SkillParam("title", "string", "Заголовок документа"),
            SkillParam("content", "string", "Текст документа"),
        ],
    )

    def _run(self, context: SkillContext, filename: str, title: str, content: str) -> SkillResult:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError:
            return SkillResult(ok=False, error="reportlab не установлен (extras: docs)")

        target = context.resolve_path(_ensure_suffix(filename, ".pdf"), subdir=DOCUMENTS_SUBDIR)
        context.policy.enforce("files.write", path=target)
        target.parent.mkdir(parents=True, exist_ok=True)

        styles = getSampleStyleSheet()
        story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
        for paragraph in content.split("\n\n"):
            if paragraph.strip():
                story.append(Paragraph(paragraph.strip(), styles["Normal"]))
                story.append(Spacer(1, 8))

        doc = SimpleDocTemplate(str(target), pagesize=A4)
        doc.build(story)
        return SkillResult(ok=True, output=f"PDF создан: {target}", data={"path": str(target)})


class CreateXlsxSkill(Skill):
    spec = SkillSpec(
        name="documents.create_xlsx",
        description="Создать таблицу Excel (.xlsx) из строк данных.",
        parameters=[
            SkillParam("filename", "string", "Имя файла, например data.xlsx"),
            SkillParam("sheet_name", "string", "Название листа", required=False),
            SkillParam(
                "rows",
                "array",
                "Строки таблицы: [[значение, ...], ...], первая строка — заголовок",
                items={"type": "array", "items": {"type": "string"}},
            ),
        ],
    )

    def _run(
        self,
        context: SkillContext,
        filename: str,
        rows: list[list],
        sheet_name: str = "Sheet1",
    ) -> SkillResult:
        try:
            from openpyxl import Workbook
        except ImportError:
            return SkillResult(ok=False, error="openpyxl не установлен (extras: docs)")

        target = context.resolve_path(_ensure_suffix(filename, ".xlsx"), subdir=DOCUMENTS_SUBDIR)
        context.policy.enforce("files.write", path=target)
        target.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        for row in rows:
            ws.append(row)
        wb.save(target)
        return SkillResult(
            ok=True,
            output=f"Таблица создана: {target} ({len(rows)} строк)",
            data={"path": str(target)},
        )


def register_document_skills(registry) -> None:
    registry.register(CreateDocxSkill())
    registry.register(CreatePptxSkill())
    registry.register(CreatePdfSkill())
    registry.register(CreateXlsxSkill())
