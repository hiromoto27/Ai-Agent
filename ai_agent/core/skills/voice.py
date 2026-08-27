"""Навыки диктофона: запись/распознавание речи, задачи, протоколы, напоминания.

Зависимости на аудио/STT (sounddevice, webrtcvad, faster-whisper) —
опциональные (extras: voice). Сам сервис (``SkillContext.voice``) собирается
в ``ai_agent/app.py``; если он не инициализирован, навыки возвращают понятную
ошибку вместо падения агента — как и остальные навыки с опциональными
зависимостями (ср. ``documents.py``, ``models.py``).
"""

from __future__ import annotations

from .base import Skill, SkillContext, SkillParam, SkillResult, SkillSpec
from .documents import DOCUMENTS_SUBDIR, _ensure_suffix


def _require_voice(context: SkillContext):
    if context.voice is None:
        raise RuntimeError(
            "сервис диктофона не инициализирован (проверьте extras 'voice' и настройки)"
        )
    return context.voice


class VoiceRecordOnceSkill(Skill):
    spec = SkillSpec(
        name="voice.record_once",
        description=(
            "Записать разовую реплику(и) с микрофона и распознать речь "
            "(останавливается по паузе в речи или по лимиту времени)."
        ),
        parameters=[
            SkillParam("max_seconds", "number", "Максимальная длительность записи в секундах", required=False),
        ],
    )

    def _run(self, context: SkillContext, max_seconds: float = 30.0) -> SkillResult:
        context.policy.enforce("mic.record")
        voice = _require_voice(context)
        utterances = voice.record_once(max_seconds=max_seconds)
        if not utterances:
            return SkillResult(ok=True, output="Речь не распознана (тишина или пусто).")
        lines = [
            f"[{u.id}] {u.text}" + (f" → задача #{u.task_id}" if u.task_id else " → не отнесено")
            for u in utterances
        ]
        return SkillResult(ok=True, output="\n".join(lines), data={"utterance_ids": [u.id for u in utterances]})


class VoiceStartContinuousSkill(Skill):
    spec = SkillSpec(
        name="voice.start_continuous",
        description=(
            "Включить постоянную (фоновую) запись микрофона с автоматической расшифровкой "
            "и классификацией реплик по задачам. Требует отдельного подтверждения политики."
        ),
        parameters=[],
    )

    def _run(self, context: SkillContext) -> SkillResult:
        context.policy.enforce("mic.continuous_start")
        voice = _require_voice(context)
        if voice.is_continuous_active:
            return SkillResult(ok=True, output="Постоянная запись уже включена.")
        voice.start_continuous()
        return SkillResult(
            ok=True,
            output="Постоянная запись включена. Не забудьте выключить её явно (voice.stop_continuous).",
        )


class VoiceStopContinuousSkill(Skill):
    spec = SkillSpec(
        name="voice.stop_continuous",
        description="Выключить постоянную (фоновую) запись микрофона.",
        parameters=[],
    )

    def _run(self, context: SkillContext) -> SkillResult:
        voice = _require_voice(context)
        if not voice.is_continuous_active:
            return SkillResult(ok=True, output="Постоянная запись не была включена.")
        voice.stop_continuous()
        return SkillResult(ok=True, output="Постоянная запись выключена.")


class VoiceTranscribeFileSkill(Skill):
    spec = SkillSpec(
        name="voice.transcribe_file",
        description="Распознать речь в существующем аудиофайле локальным STT-движком.",
        parameters=[SkillParam("path", "string", "Путь к аудиофайлу (относительно workspace или абсолютный)")],
    )

    def _run(self, context: SkillContext, path: str) -> SkillResult:
        voice = _require_voice(context)
        target = context.resolve_path(path)
        context.policy.enforce("files.read", path=target)
        if not target.exists():
            return SkillResult(ok=False, error=f"файл не найден: {target}")
        text = voice.transcribe_file(target)
        return SkillResult(ok=True, output=text or "(речь не распознана)", data={"text": text})


class VoiceConfirmTaskSkill(Skill):
    spec = SkillSpec(
        name="voice.confirm_task",
        description=(
            "Подтвердить или отклонить принадлежность реплики к задаче — "
            "обучающий сигнал, повышающий точность будущей классификации для этой задачи."
        ),
        parameters=[
            SkillParam("utterance_id", "integer", "ID реплики"),
            SkillParam("task_id", "integer", "ID задачи"),
            SkillParam("belongs", "boolean", "true — реплика относится к задаче, false — не относится"),
        ],
    )

    def _run(self, context: SkillContext, utterance_id: int, task_id: int, belongs: bool) -> SkillResult:
        voice = _require_voice(context)
        if voice.store.get_task(task_id) is None:
            return SkillResult(ok=False, error=f"задача #{task_id} не найдена")
        if voice.store.get_utterance(utterance_id) is None:
            return SkillResult(ok=False, error=f"реплика #{utterance_id} не найдена")
        if belongs:
            voice.confirm_task(utterance_id, task_id)
            return SkillResult(ok=True, output=f"Принято: реплика #{utterance_id} отнесена к задаче #{task_id}.")
        voice.reject_task(utterance_id, task_id)
        return SkillResult(ok=True, output=f"Принято: реплика #{utterance_id} не относится к задаче #{task_id}.")


class VoiceGenerateProtocolSkill(Skill):
    spec = SkillSpec(
        name="voice.generate_protocol",
        description="Собрать протокол разговора(ов) по задачам и сохранить как документ Word (.docx).",
        parameters=[
            SkillParam("filename", "string", "Имя файла протокола, например protocol.docx", required=False),
            SkillParam(
                "recording_id",
                "integer",
                "Ограничить одной записанной сессией (иначе — все накопленные реплики)",
                required=False,
            ),
        ],
    )

    def _run(self, context: SkillContext, filename: str = "protocol.docx", recording_id: int | None = None) -> SkillResult:
        voice = _require_voice(context)
        protocol = voice.generate_protocol(recording_id=recording_id)

        try:
            from docx import Document
        except ImportError:
            return SkillResult(ok=False, error="python-docx не установлен (extras: docs)")

        target = context.resolve_path(_ensure_suffix(filename, ".docx"), subdir=DOCUMENTS_SUBDIR)
        context.policy.enforce("files.write", path=target)
        target.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()
        doc.add_heading("Протокол разговора", level=0)
        for section in protocol.sections:
            doc.add_heading(section.task_title, level=1)
            for item in section.items:
                doc.add_paragraph(item, style="List Bullet")
            if section.action_items:
                doc.add_heading("Пункты-действия", level=2)
                for action in section.action_items:
                    doc.add_paragraph(action, style="List Number")
        if protocol.unassigned:
            doc.add_heading("Не отнесено к задачам", level=1)
            for item in protocol.unassigned:
                doc.add_paragraph(item, style="List Bullet")
        doc.save(target)

        from ai_agent.core.voice.protocol import protocol_to_text

        return SkillResult(
            ok=True,
            output=f"Протокол создан: {target}\n\n{protocol_to_text(protocol)}",
            data={"path": str(target)},
        )


class TasksCreateSkill(Skill):
    spec = SkillSpec(
        name="tasks.create",
        description=(
            "Создать задачу для диктофона: реплики диктофона будут классифицироваться "
            "по контекстным словам как относящиеся к ней или нет."
        ),
        parameters=[
            SkillParam("title", "string", "Название задачи"),
            SkillParam(
                "description",
                "string",
                "Описание/контекст задачи (используется для первичной классификации)",
                required=False,
            ),
            SkillParam("due_at", "number", "Срок задачи (unix timestamp), опционально", required=False),
        ],
    )

    def _run(self, context: SkillContext, title: str, description: str = "", due_at: float | None = None) -> SkillResult:
        voice = _require_voice(context)
        task = voice.create_task(title=title, description=description, due_at=due_at)
        return SkillResult(ok=True, output=f"Задача создана: #{task.id} «{task.title}»", data={"task_id": task.id})


class TasksListSkill(Skill):
    spec = SkillSpec(
        name="tasks.list",
        description="Показать список задач диктофона с уровнем уверенности классификатора по каждой.",
        parameters=[],
    )

    def _run(self, context: SkillContext) -> SkillResult:
        voice = _require_voice(context)
        tasks = voice.store.list_tasks()
        if not tasks:
            return SkillResult(ok=True, output="Задач пока нет.")
        lines = []
        for task in tasks:
            profile = voice.classifier.profiles.get(task.id)
            confidence = f"{profile.confidence:.0%}" if profile else "—"
            lines.append(f"#{task.id} «{task.title}» — уверенность классификатора: {confidence}")
        return SkillResult(ok=True, output="\n".join(lines))


class TasksSetReminderSkill(Skill):
    spec = SkillSpec(
        name="tasks.set_reminder",
        description="Поставить напоминание по задаче на определённое время.",
        parameters=[
            SkillParam("task_id", "integer", "ID задачи"),
            SkillParam("fire_at", "number", "Время напоминания (unix timestamp)"),
            SkillParam("message", "string", "Текст напоминания", required=False),
        ],
    )

    def _run(self, context: SkillContext, task_id: int, fire_at: float, message: str = "") -> SkillResult:
        voice = _require_voice(context)
        task = voice.store.get_task(task_id)
        if task is None:
            return SkillResult(ok=False, error=f"задача #{task_id} не найдена")
        reminder = voice.store.add_reminder(
            task_id=task_id, fire_at=fire_at, message=message or f"Напоминание по задаче «{task.title}»"
        )
        return SkillResult(
            ok=True,
            output=f"Напоминание #{reminder.id} поставлено на {reminder.fire_at}",
            data={"reminder_id": reminder.id},
        )


def register_voice_skills(registry) -> None:
    registry.register(VoiceRecordOnceSkill())
    registry.register(VoiceStartContinuousSkill())
    registry.register(VoiceStopContinuousSkill())
    registry.register(VoiceTranscribeFileSkill())
    registry.register(VoiceConfirmTaskSkill())
    registry.register(VoiceGenerateProtocolSkill())
    registry.register(TasksCreateSkill())
    registry.register(TasksListSkill())
    registry.register(TasksSetReminderSkill())
