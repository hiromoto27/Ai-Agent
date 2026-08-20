from ai_agent.core.skills.documents import (
    CreateDocxSkill,
    CreatePdfSkill,
    CreatePptxSkill,
    CreateXlsxSkill,
)


def test_create_docx(permissive_context):
    result = CreateDocxSkill().run(
        permissive_context,
        filename="report",
        title="Отчёт",
        content="Первый абзац.\n\nВторой абзац.",
    )
    assert result.ok
    path = permissive_context.workspace_root / "documents" / "report.docx"
    assert path.exists()

    from docx import Document

    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Отчёт" in text
    assert "Первый абзац." in text


def test_create_pptx(permissive_context):
    result = CreatePptxSkill().run(
        permissive_context,
        filename="deck.pptx",
        title="Презентация",
        slides=[
            {"title": "Слайд 1", "bullets": ["Пункт A", "Пункт B"]},
            {"title": "Слайд 2", "bullets": ["Пункт C"]},
        ],
    )
    assert result.ok
    path = permissive_context.workspace_root / "documents" / "deck.pptx"
    assert path.exists()

    from pptx import Presentation

    prs = Presentation(path)
    assert len(prs.slides) == 3  # титульный + 2


def test_create_pdf(permissive_context):
    result = CreatePdfSkill().run(
        permissive_context, filename="report.pdf", title="Отчёт", content="Текст документа."
    )
    assert result.ok
    path = permissive_context.workspace_root / "documents" / "report.pdf"
    assert path.exists()
    assert path.read_bytes().startswith(b"%PDF")


def test_create_xlsx(permissive_context):
    result = CreateXlsxSkill().run(
        permissive_context,
        filename="data.xlsx",
        rows=[["Имя", "Возраст"], ["Аня", "30"], ["Боб", "25"]],
    )
    assert result.ok
    path = permissive_context.workspace_root / "documents" / "data.xlsx"
    assert path.exists()

    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    assert ws["A1"].value == "Имя"
    assert ws["B2"].value == "30"


def test_create_docx_outside_workspace_denied_when_locked(locked_context, tmp_path):
    outside = tmp_path / "outside" / "report.docx"
    result = CreateDocxSkill().run(
        locked_context, filename=str(outside), title="X", content="Y"
    )
    assert not result.ok
    assert not outside.exists()
